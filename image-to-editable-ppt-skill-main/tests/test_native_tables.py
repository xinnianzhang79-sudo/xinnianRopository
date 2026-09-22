import copy
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image

RUNTIME = Path(__file__).resolve().parents[1] / "skills/image-to-editable-ppt/cli/editppt/runtime"
sys.path.insert(0, str(RUNTIME))
from build_pptx_from_manifest import normalize_manifest, render_preview, write_pptx, write_deck

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def fixture():
    return {"slide": {"width": 8, "height": 4.5}, "preview_scale": 100,
            "tables": [{"box": [1, 1, 6, 2], "column_widths": [1, 2, 1],
                        "style": {"font_size": 12, "fill": "#ABCDEF", "fit_text": False},
                        "cells": [[{"text": "Merged & editable", "row_span": 2, "col_span": 2}, "", "A"],
                                  ["", "", {"text": "B", "style": {"bold": True, "align": "right"}}],
                                  ["C", "D", "E"]]}]}


class NativeTablesTest(unittest.TestCase):
    def test_native_merge_and_editable_cell_structure(self):
        data = fixture()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "manifest.json"
            path.write_text(json.dumps(data))
            ppt = Path(directory) / "test.pptx"
            write_pptx(data, ppt, path)
            with zipfile.ZipFile(ppt) as archive:
                root = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
            self.assertEqual(1, len(root.findall(".//a:tbl", NS)))
            self.assertEqual(0, len(root.findall(".//p:sp", NS)))
            rows = root.findall(".//a:tr", NS)
            self.assertEqual({"rowSpan": "2", "gridSpan": "2"}, rows[0][0].attrib)
            self.assertEqual({"rowSpan": "2", "hMerge": "1"}, rows[0][1].attrib)
            self.assertEqual({"gridSpan": "2", "vMerge": "1"}, rows[1][0].attrib)
            self.assertEqual({"hMerge": "1", "vMerge": "1"}, rows[1][1].attrib)
            texts = [node.text for node in root.findall(".//a:t", NS) if node.text]
            self.assertEqual(["Merged & editable", "A", "B", "C", "D", "E"], texts)
            self.assertEqual("ABCDEF", root.find(".//a:tcPr/a:solidFill/a:srgbClr", NS).get("val"))
            try:
                from pptx import Presentation
            except ImportError:
                return  # Optional independent reader; XML assertions above always run.
            opened = Presentation(ppt)
            table = opened.slides[0].shapes[0].table
            self.assertTrue(table.cell(0, 0).is_merge_origin)
            self.assertEqual((2, 2), (table.cell(0, 0).span_width, table.cell(0, 0).span_height))
            self.assertTrue(table.cell(1, 1).is_spanned)
            table.cell(2, 1).text = "edited"
            opened.save(ppt)
            self.assertEqual("edited", Presentation(ppt).slides[0].shapes[0].table.cell(2, 1).text)

    def test_pixel_position_weights_and_no_input_mutation(self):
        data = fixture()
        table = data["tables"][0]
        del table["box"]
        table["box_px"] = [100, 50, 600, 200]
        data["source"] = {"width_px": 800, "height_px": 450}
        before = copy.deepcopy(data)
        norm = normalize_manifest(data)["tables"][0]
        self.assertEqual(before, data)
        normalized = normalize_manifest(data)
        self.assertEqual(normalized, normalize_manifest(normalized))
        self.assertEqual((1, .5, 6, 2), tuple(norm[k] for k in ("left", "top", "width", "height")))
        self.assertEqual([1.5, 3, 1.5], norm["column_widths"])
        self.assertEqual(4.5, norm["cells"][0][0]["width"])

    def test_invalid_tables_fail_before_writing(self):
        changes = [
            {"cells": []}, {"cells": [["a"], ["b", "c"]]},
            {"column_widths": [1, 2]}, {"row_heights": [1, 0, 1]},
            {"row_heights": [1, float("nan"), 1]}, {"box": [0, 0, -1, 2]},
            {"style": {"margin_left": 100}},
            {"cells": [[{"text": "a", "col_span": 2}, "covered text"]]},
            {"cells": [[{"text": "a", "row_span": 2}]]},
            {"cells": [[{"text": "a", "col_span": True}]]},
            {"cells": [[None]]}, {"cells": [[{"text": 42}]]},
            {"cells": [["a", {"row_span": 2}], [{"col_span": 2}, ""]]},
        ]
        for change in changes:
            with self.subTest(change=change):
                data = fixture()
                data["tables"][0].update(change)
                with self.assertRaises(ValueError):
                    normalize_manifest(data)

    def test_unsupported_cell_content_and_style_are_rejected(self):
        for key in ("runs", "paragraphs", "image", "images", "rotation"):
            for target in ("cell", "cell_style", "table_style"):
                with self.subTest(key=key, target=target):
                    data = fixture()
                    table = data["tables"][0]
                    cell = {"text": "preserve this text"}
                    table["cells"] = [[cell]]
                    table.pop("column_widths")
                    destination = cell if target == "cell" else cell.setdefault("style", {}) if target == "cell_style" else table["style"]
                    destination[key] = [{"text": "must not replace or erase cell text"}]
                    with self.assertRaisesRegex(ValueError, "unsupported fields"):
                        normalize_manifest(data)

    def test_preview_merge_and_layering(self):
        data = fixture()
        data["tables"][0]["cells"][0][0]["text"] = ""
        data["shapes"] = [{"type": "rect", "left": 0, "top": 0, "width": 8, "height": 4.5, "fill": "#FF0000", "z_index": 100}]
        with tempfile.TemporaryDirectory() as directory:
            out = Path(directory) / "preview.png"
            render_preview(data, Path(directory) / "manifest.json", out)
            with Image.open(out) as image:
                self.assertEqual((171, 205, 239), image.getpixel((250, 160)))
                self.assertEqual((255, 0, 0), image.getpixel((20, 20)))
            data["shapes"][0]["z_index"] = 999
            render_preview(data, Path(directory) / "manifest.json", out)
            with Image.open(out) as image:
                self.assertEqual((255, 0, 0), image.getpixel((250, 160)))

    def test_cell_text_fit_uses_inner_box_and_xml_layer_order(self):
        from build_pptx_from_manifest import slide_xml
        data = fixture()
        table = data["tables"][0]
        table["style"].update(fit_text=True, font_size=72)
        table["cells"] = [["A very long editable cell requiring smaller text"]]
        table.pop("column_widths")
        normalized = normalize_manifest(data)
        self.assertLess(normalized["tables"][0]["cells"][0][0]["font_size"], 72)
        self.assertEqual(normalized, normalize_manifest(normalized))
        data["text_boxes"] = [{"text": "front", "left": 0, "top": 0, "width": 1, "height": 1}]
        root = ET.fromstring(slide_xml(normalize_manifest(data)))
        tree = root.find(".//p:spTree", NS)
        self.assertTrue(tree[-2].tag.endswith("}graphicFrame"))
        self.assertTrue(tree[-1].tag.endswith("}sp"))

    def test_legacy_and_deck_build(self):
        legacy = {"text_boxes": [{"text": "ordinary", "left": 1, "top": 1, "width": 2, "height": 1}]}
        with tempfile.TemporaryDirectory() as directory:
            out = Path(directory) / "deck.pptx"
            path = Path(directory) / "manifest.json"
            write_deck({}, [{"manifest": legacy, "manifest_path": path}, {"manifest": fixture(), "manifest_path": path}], out, [])
            with zipfile.ZipFile(out) as archive:
                first = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
                second = ET.fromstring(archive.read("ppt/slides/slide2.xml"))
            self.assertEqual(0, len(first.findall(".//a:tbl", NS)))
            self.assertEqual(1, len(second.findall(".//a:tbl", NS)))


if __name__ == "__main__":
    unittest.main()
